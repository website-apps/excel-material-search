from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.database import connect_database


class ExcelIndexError(ValueError):
    """Raised when an uploaded workbook cannot be indexed."""


class ExcelIndexStore:
    SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
    EXCLUDED_HEADER_KEYS = {
        "reference",
        "referencedesignator",
        "refdes",
        "designator",
        "位号",
        "参考标号",
        "参考位号",
    }
    HEADER_SCAN_ROWS = 100

    def __init__(self, db_file: Path, storage_dir: Path) -> None:
        self.db_file = db_file
        self.storage_dir = storage_dir
        self.initialize()

    def initialize(self) -> None:
        self.db_file.parent.mkdir(parents=True, exist_ok=True)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS excel_files (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    original_name TEXT NOT NULL,
                    stored_name TEXT NOT NULL,
                    stored_path TEXT NOT NULL,
                    sha256 TEXT NOT NULL UNIQUE,
                    size_bytes INTEGER NOT NULL,
                    status TEXT NOT NULL,
                    sheet_count INTEGER NOT NULL DEFAULT 0,
                    indexed_cell_count INTEGER NOT NULL DEFAULT 0,
                    error TEXT,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS excel_sheets (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_id INTEGER NOT NULL REFERENCES excel_files(id) ON DELETE CASCADE,
                    sheet_name TEXT NOT NULL,
                    sheet_index INTEGER NOT NULL,
                    header_row INTEGER,
                    max_row INTEGER,
                    max_column INTEGER,
                    UNIQUE(file_id, sheet_name)
                );

                CREATE TABLE IF NOT EXISTS excel_cells (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_id INTEGER NOT NULL REFERENCES excel_files(id) ON DELETE CASCADE,
                    sheet_id INTEGER NOT NULL REFERENCES excel_sheets(id) ON DELETE CASCADE,
                    sheet_name TEXT NOT NULL,
                    row_number INTEGER NOT NULL,
                    column_number INTEGER NOT NULL,
                    cell_address TEXT NOT NULL,
                    field_name TEXT,
                    raw_value TEXT NOT NULL,
                    normalized_value TEXT NOT NULL,
                    row_context_json TEXT NOT NULL
                );

                CREATE INDEX IF NOT EXISTS idx_excel_cells_normalized_value
                    ON excel_cells(normalized_value);
                CREATE INDEX IF NOT EXISTS idx_excel_cells_file_id
                    ON excel_cells(file_id);
                """
            )

    def index_workbook(self, original_name: str, payload: bytes) -> dict[str, Any]:
        suffix = Path(original_name).suffix.lower()
        if suffix not in self.SUPPORTED_SUFFIXES:
            raise ExcelIndexError("unsupported Excel type; use .xlsx or .xlsm")
        if not payload:
            raise ExcelIndexError("Excel file is empty")

        digest = hashlib.sha256(payload).hexdigest()
        existing = self._get_file_by_hash(digest)
        if existing is not None:
            return existing

        stored_name = f"{uuid4().hex}{suffix}"
        stored_path = self.storage_dir / stored_name
        created_at = datetime.now(timezone.utc).isoformat()

        try:
            workbook = load_workbook(
                io.BytesIO(payload),
                data_only=True,
                read_only=True,
            )
        except Exception as exc:  # openpyxl raises several format-specific errors
            raise ExcelIndexError(f"unable to read workbook: {exc}") from exc

        file_id: int | None = None
        sheet_count = 0
        indexed_cell_count = 0
        try:
            with self._connect() as connection:
                cursor = connection.execute(
                    """
                    INSERT INTO excel_files (
                        original_name, stored_name, stored_path, sha256, size_bytes,
                        status, created_at
                    ) VALUES (?, ?, ?, ?, ?, 'indexing', ?)
                    """,
                    (
                        original_name,
                        stored_name,
                        str(stored_path),
                        digest,
                        len(payload),
                        created_at,
                    ),
                )
                file_id = int(cursor.lastrowid)

                for sheet_index, worksheet in enumerate(workbook.worksheets):
                    sheet_count += 1
                    header_row, headers = self._detect_headers(worksheet)
                    sheet_cursor = connection.execute(
                        """
                        INSERT INTO excel_sheets (
                            file_id, sheet_name, sheet_index, header_row, max_row, max_column
                        ) VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            file_id,
                            worksheet.title,
                            sheet_index,
                            header_row,
                            worksheet.max_row or 0,
                            worksheet.max_column or 0,
                        ),
                    )
                    sheet_id = int(sheet_cursor.lastrowid)

                    for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                        row_values = {
                            cell.column: self._display_value(cell.value)
                            for cell in row
                            if cell.value is not None
                        }
                        if not row_values or row_number == header_row:
                            continue

                        row_context = {
                            get_column_letter(column): value
                            for column, value in row_values.items()
                        }
                        row_context_json = json.dumps(
                            row_context,
                            ensure_ascii=False,
                            separators=(",", ":"),
                        )

                        for cell in row:
                            if cell.value is None:
                                continue
                            field_name = headers.get(cell.column)
                            if self._is_excluded_field(field_name):
                                continue
                            raw_value = self._display_value(cell.value)
                            normalized_value = normalize_excel_value(raw_value)
                            if not normalized_value:
                                continue

                            connection.execute(
                                """
                                INSERT INTO excel_cells (
                                    file_id, sheet_id, sheet_name, row_number, column_number,
                                    cell_address, field_name, raw_value, normalized_value,
                                    row_context_json
                                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                """,
                                (
                                    file_id,
                                    sheet_id,
                                    worksheet.title,
                                    row_number,
                                    cell.column,
                                    cell.coordinate,
                                    field_name,
                                    raw_value,
                                    normalized_value,
                                    row_context_json,
                                ),
                            )
                            indexed_cell_count += 1

                connection.execute(
                    """
                    UPDATE excel_files
                    SET status = 'ready', sheet_count = ?, indexed_cell_count = ?
                    WHERE id = ?
                    """,
                    (sheet_count, indexed_cell_count, file_id),
                )
            stored_path.write_bytes(payload)
        except Exception as exc:
            if file_id is not None:
                with self._connect() as connection:
                    connection.execute(
                        "UPDATE excel_files SET status = 'failed', error = ? WHERE id = ?",
                        (str(exc), file_id),
                    )
            raise ExcelIndexError(f"failed to index workbook: {exc}") from exc
        finally:
            workbook.close()

        result = self._get_file(file_id)
        if result is None:
            raise ExcelIndexError("indexed workbook metadata is missing")
        return result

    def search(self, query: str, limit: int = 200) -> list[dict[str, Any]]:
        normalized_query = normalize_excel_value(query)
        if not normalized_query:
            return []

        select_query = """
            SELECT
                excel_cells.file_id,
                excel_files.original_name,
                excel_files.created_at,
                excel_cells.sheet_name,
                excel_cells.row_number,
                excel_cells.cell_address,
                excel_cells.field_name,
                excel_cells.raw_value,
                excel_cells.row_context_json
            FROM excel_cells
            JOIN excel_files ON excel_files.id = excel_cells.file_id
            WHERE excel_files.status = 'ready'
        """
        with self._connect() as connection:
            rows = connection.execute(
                f"""{select_query}
                  AND excel_cells.normalized_value = ?
                ORDER BY excel_files.original_name, excel_cells.sheet_name,
                         excel_cells.row_number, excel_cells.column_number
                LIMIT ?
                """,
                (normalized_query, limit),
            ).fetchall()
            if not rows:
                escaped_query = _escape_like_query(normalized_query)
                rows = connection.execute(
                    f"""{select_query}
                      AND excel_cells.normalized_value LIKE ? ESCAPE '\\'
                    ORDER BY excel_files.original_name, excel_cells.sheet_name,
                             excel_cells.row_number, excel_cells.column_number
                    LIMIT ?
                    """,
                    (f"%{escaped_query}%", limit),
                ).fetchall()

        return [
            {
                "file_id": row["file_id"],
                "file_name": row["original_name"],
                "created_at": row["created_at"],
                "sheet_name": row["sheet_name"],
                "row_number": row["row_number"],
                "cell_address": row["cell_address"],
                "field_name": row["field_name"],
                "raw_value": row["raw_value"],
                "row_context": json.loads(row["row_context_json"]),
            }
            for row in rows
        ]

    def list_files(self) -> list[dict[str, Any]]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT id, original_name, status, sheet_count, indexed_cell_count,
                       error, created_at
                FROM excel_files
                ORDER BY created_at DESC, id DESC
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def delete_file(self, file_id: int) -> dict[str, Any] | None:
        deleted_files = self.delete_files([file_id])
        return deleted_files[0] if deleted_files else None

    def delete_files(self, file_ids: list[int]) -> list[dict[str, Any]] | None:
        unique_file_ids = sorted(set(file_ids))
        if not unique_file_ids:
            return []

        placeholders = ",".join("?" for _ in unique_file_ids)
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT id, original_name, stored_path, status, sheet_count,
                       indexed_cell_count, error, created_at, sha256, size_bytes
                FROM excel_files
                WHERE id IN ({placeholders})
                ORDER BY id
                """,
                tuple(unique_file_ids),
            ).fetchall()
            if len(rows) != len(unique_file_ids):
                return None

            deleted_files = [dict(row) for row in rows]
            connection.execute(
                f"DELETE FROM excel_files WHERE id IN ({placeholders})",
                tuple(unique_file_ids),
            )

        for deleted_file in deleted_files:
            try:
                Path(deleted_file["stored_path"]).unlink(missing_ok=True)
            except OSError:
                # Database deletion is authoritative; an inaccessible orphan can be cleaned up later.
                pass
        return deleted_files

    def _get_file_by_hash(self, digest: str) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT id FROM excel_files WHERE sha256 = ?",
                (digest,),
            ).fetchone()
        return self._get_file(row["id"]) if row is not None else None

    def _get_file(self, file_id: int | None) -> dict[str, Any] | None:
        if file_id is None:
            return None
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT id, original_name, stored_path, status, sheet_count,
                       indexed_cell_count, error, created_at, sha256, size_bytes
                FROM excel_files
                WHERE id = ?
                """,
                (file_id,),
            ).fetchone()
        return dict(row) if row is not None else None

    def _detect_headers(self, worksheet: Any) -> tuple[int | None, dict[int, str]]:
        candidates: list[tuple[int, int, dict[int, str]]] = []
        max_row = min(worksheet.max_row or 0, self.HEADER_SCAN_ROWS)
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=max_row), start=1
        ):
            values = {
                cell.column: self._display_value(cell.value)
                for cell in row
                if cell.value is not None and self._display_value(cell.value)
            }
            if len(values) < 2:
                continue
            alias_score = sum(1 for value in values.values() if self._looks_like_header(value))
            text_score = sum(1 for value in values.values() if not self._looks_numeric(value))
            score = len(values) * 3 + alias_score * 8 + text_score
            candidates.append((score, row_number, values))

        if not candidates:
            return None, {}
        _, header_row, values = max(candidates, key=lambda candidate: (candidate[0], -candidate[1]))
        return header_row, values

    @classmethod
    def _looks_like_header(cls, value: str) -> bool:
        key = _header_key(value)
        return key in cls.EXCLUDED_HEADER_KEYS or key in {
            "partnumber",
            "mpn",
            "物料号",
            "料号",
            "型号",
            "description",
            "package",
            "value",
        }

    @staticmethod
    def _looks_numeric(value: str) -> bool:
        return bool(re.fullmatch(r"[+-]?(?:\d+(?:\.\d+)?|\.\d+)", value.strip()))

    @classmethod
    def _is_excluded_field(cls, field_name: str | None) -> bool:
        return bool(field_name and _header_key(field_name) in cls.EXCLUDED_HEADER_KEYS)

    @staticmethod
    def _display_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        return str(value).strip()

    def _connect(self) -> sqlite3.Connection:
        connection = connect_database(self.db_file, "archive")
        connection.row_factory = sqlite3.Row
        return connection


def _header_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).strip().lower()
    return re.sub(r"[\s_\-/:：()（）]+", "", normalized)


def normalize_excel_value(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    return re.sub(r"\s+", " ", normalized)


def _escape_like_query(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
