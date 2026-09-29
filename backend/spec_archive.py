from __future__ import annotations

import io
import re
import sqlite3
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.database import connect_database
from backend.excel_index import ExcelIndexStore


class ArchiveError(ValueError):
    """Raised when a document cannot be stored or indexed."""


class ArchiveStore:
    MANUAL_SUFFIXES = {".pdf", ".docx", ".xls", ".xlsx", ".xlsm", ".xltx", ".xltm", ".txt", ".md"}
    MAX_INDEX_CHARS = 500_000
    MAX_PREVIEW_ROWS = 200
    MAX_PREVIEW_COLUMNS = 40

    def __init__(self, db_file: Path, storage_dir: Path, excel_index: ExcelIndexStore) -> None:
        self.db_file = db_file
        self.storage_dir = storage_dir
        self.excel_index = excel_index
        self.initialize()

    def initialize(self) -> None:
        self.db_file.parent.mkdir(parents=True, exist_ok=True)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS archive_documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    excel_file_id INTEGER UNIQUE REFERENCES excel_files(id) ON DELETE CASCADE,
                    kind TEXT NOT NULL CHECK (kind IN ('manual', 'bom')),
                    title TEXT NOT NULL,
                    category TEXT NOT NULL DEFAULT '',
                    package TEXT NOT NULL DEFAULT '',
                    vendor TEXT NOT NULL DEFAULT '',
                    remark TEXT NOT NULL DEFAULT '',
                    note TEXT NOT NULL DEFAULT '',
                    board_code TEXT NOT NULL DEFAULT '',
                    board_type TEXT NOT NULL DEFAULT '',
                    main_chip TEXT NOT NULL DEFAULT '',
                    board_name TEXT NOT NULL DEFAULT '',
                    original_name TEXT NOT NULL,
                    stored_name TEXT NOT NULL,
                    sha256 TEXT NOT NULL UNIQUE,
                    size_bytes INTEGER NOT NULL DEFAULT 0,
                    extension TEXT NOT NULL,
                    mime_type TEXT NOT NULL,
                    content_text TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'ready',
                    error TEXT,
                    created_at TEXT NOT NULL,
                    file_modified_at TEXT NOT NULL DEFAULT ''
                );

                CREATE INDEX IF NOT EXISTS idx_archive_documents_kind
                    ON archive_documents(kind);
                CREATE INDEX IF NOT EXISTS idx_archive_documents_created_at
                    ON archive_documents(created_at);

                CREATE VIRTUAL TABLE IF NOT EXISTS archive_documents_fts
                    USING fts5(document_id UNINDEXED, content, tokenize='trigram');
                """
            )
        self._migrate_existing_excel_files()

    def list_documents(
        self,
        *,
        query: str = "",
        kind: str = "",
        category: str = "",
        vendor: str = "",
        board_code: str = "",
        main_chip: str = "",
    ) -> list[dict[str, Any]]:
        rows = self._load_rows()
        return [
            self._public_document(row)
            for row in rows
            if self._matches(row, query, kind, category, vendor, board_code, main_chip)
        ]

    def query_documents(self, query: str, *, kind: str = "bom") -> dict[str, Any]:
        normalized = _normalize_text(query)
        if not normalized:
            raise ArchiveError("query is required")
        if kind == "bom":
            return {"query": query.strip(), "matches": self.excel_index.search(query)}
        if kind != "manual":
            raise ArchiveError("kind must be manual or bom")

        rows = self._load_rows()
        matches = []
        terms = _query_terms(normalized)
        for row in rows:
            if row["kind"] != "manual" or row["status"] not in {"ready", "empty"}:
                continue
            haystack = _normalize_text(self._search_text(row))
            if terms and not all(term in haystack for term in terms):
                continue
            excerpt = _excerpt(row["content_text"], terms) or row["note"] or row["remark"]
            matches.append(
                {
                    "document_id": row["id"],
                    "title": row["title"],
                    "file_name": row["original_name"],
                    "category": row["category"] or "未分类",
                    "vendor": row["vendor"] or "",
                    "excerpt": excerpt,
                    "created_at": row["created_at"],
                }
            )
        return {"query": query.strip(), "matches": matches[:200]}

    def upload_document(
        self,
        original_name: str,
        payload: bytes,
        *,
        kind: str,
        metadata: dict[str, Any] | None = None,
        file_modified_at: str = "",
    ) -> dict[str, Any]:
        safe_name = Path(original_name or "").name
        suffix = Path(safe_name).suffix.lower()
        metadata = metadata or {}
        if not safe_name or suffix not in self.MANUAL_SUFFIXES:
            raise ArchiveError("unsupported document type")
        if kind not in {"manual", "bom"}:
            raise ArchiveError("kind must be manual or bom")
        if not payload:
            raise ArchiveError("document is empty")

        digest = _sha256(payload)
        existing = self._get_row_by_hash(digest)
        if existing is not None:
            return self._public_document(existing)

        if kind == "bom":
            if suffix not in ExcelIndexStore.SUPPORTED_SUFFIXES:
                raise ArchiveError("BOM requires an Excel workbook")
            indexed = self.excel_index.index_workbook(safe_name, payload)
            return self.register_excel_file(indexed, file_modified_at=file_modified_at)

        stored_name = f"{uuid4().hex}{suffix}"
        stored_path = self.storage_dir / stored_name
        try:
            content_text = _extract_document_text(payload, suffix)
            stored_path.write_bytes(payload)
            now = datetime.now(timezone.utc).isoformat()
            title = _text(metadata.get("title")) or Path(safe_name).stem
            record = {
                "excel_file_id": None,
                "kind": "manual",
                "title": title[:160],
                "category": _text(metadata.get("category")) or "未分类",
                "package": _text(metadata.get("package")),
                "vendor": _text(metadata.get("vendor")),
                "remark": _text(metadata.get("remark")),
                "note": _text(metadata.get("note")),
                "board_code": "",
                "board_type": "",
                "main_chip": "",
                "board_name": "",
                "original_name": safe_name,
                "stored_name": stored_name,
                "sha256": digest,
                "size_bytes": len(payload),
                "extension": suffix[1:],
                "mime_type": mime_type(suffix),
                "content_text": content_text[: self.MAX_INDEX_CHARS],
                "status": "ready" if content_text else "empty",
                "error": None,
                "created_at": now,
                "file_modified_at": file_modified_at,
            }
            with self._connect() as connection:
                cursor = connection.execute(
                    """
                    INSERT INTO archive_documents (
                        excel_file_id, kind, title, category, package, vendor, remark, note,
                        board_code, board_type, main_chip, board_name, original_name,
                        stored_name, sha256, size_bytes, extension, mime_type, content_text,
                        status, error, created_at, file_modified_at
                    ) VALUES (:excel_file_id, :kind, :title, :category, :package, :vendor,
                        :remark, :note, :board_code, :board_type, :main_chip, :board_name,
                        :original_name, :stored_name, :sha256, :size_bytes, :extension,
                        :mime_type, :content_text, :status, :error, :created_at,
                        :file_modified_at)
                    """,
                    record,
                )
                document_id = int(cursor.lastrowid)
            self._replace_fts(document_id, record["content_text"] + " " + self._search_text(record))
            return self._public_document(self._get_row(document_id))
        except Exception:
            stored_path.unlink(missing_ok=True)
            raise

    def register_excel_file(
        self,
        indexed_file: dict[str, Any],
        *,
        file_modified_at: str = "",
    ) -> dict[str, Any]:
        existing = self._get_row_by_excel_id(indexed_file["id"])
        if existing is not None:
            return self._public_document(existing)
        parsed = parse_bom_filename(indexed_file["original_name"])
        record = {
            "excel_file_id": indexed_file["id"],
            "kind": "bom",
            "title": Path(indexed_file["original_name"]).stem,
            "category": "BOM",
            "package": "",
            "vendor": "",
            "remark": "",
            "note": "",
            "board_code": parsed["board_code"],
            "board_type": parsed["board_type"],
            "main_chip": parsed["main_chip"],
            "board_name": parsed["board_name"],
            "original_name": indexed_file["original_name"],
            "stored_name": Path(indexed_file["stored_path"]).name,
            "sha256": indexed_file["sha256"],
            "size_bytes": indexed_file["size_bytes"],
            "extension": Path(indexed_file["original_name"]).suffix.lower().lstrip("."),
            "mime_type": mime_type(Path(indexed_file["original_name"]).suffix.lower()),
            "content_text": self._excel_content(indexed_file["id"]),
            "status": indexed_file["status"],
            "error": indexed_file.get("error"),
            "created_at": indexed_file["created_at"],
            "file_modified_at": file_modified_at,
        }
        with self._connect() as connection:
            cursor = connection.execute(
                """
                INSERT INTO archive_documents (
                    excel_file_id, kind, title, category, package, vendor, remark, note,
                    board_code, board_type, main_chip, board_name, original_name,
                    stored_name, sha256, size_bytes, extension, mime_type, content_text,
                    status, error, created_at, file_modified_at
                ) VALUES (:excel_file_id, :kind, :title, :category, :package, :vendor,
                    :remark, :note, :board_code, :board_type, :main_chip, :board_name,
                    :original_name, :stored_name, :sha256, :size_bytes, :extension,
                    :mime_type, :content_text, :status, :error, :created_at,
                    :file_modified_at)
                """,
                record,
            )
            document_id = int(cursor.lastrowid)
        self._replace_fts(document_id, record["content_text"] + " " + self._search_text(record))
        return self._public_document(self._get_row(document_id))

    def update_document(self, document_id: int, payload: dict[str, Any]) -> dict[str, Any] | None:
        row = self._get_row(document_id)
        if row is None:
            return None
        fields = (
            ("title", 160), ("category", 80), ("package", 80), ("vendor", 80),
            ("remark", 240), ("note", 500), ("board_type", 40), ("main_chip", 80),
            ("board_name", 160),
        )
        updates = {name: _text(payload[name])[:limit] for name, limit in fields if name in payload}
        if row["kind"] == "bom":
            board_type = updates.get("board_type", row["board_type"])
            updates["board_code"] = "RD" if board_type == "开发板" else "PD" if board_type == "产品板" else ""
        if updates:
            assignments = ", ".join(f"{name} = :{name}" for name in updates)
            with self._connect() as connection:
                connection.execute(
                    f"UPDATE archive_documents SET {assignments} WHERE id = :id",
                    {**updates, "id": document_id},
                )
            self._replace_fts(document_id, self._search_text(self._get_row(document_id)))
        return self._public_document(self._get_row(document_id))

    def delete_document(self, document_id: int) -> dict[str, Any] | None:
        row = self._get_row(document_id)
        if row is None:
            return None
        if row["excel_file_id"] is not None:
            deleted = self.excel_index.delete_file(int(row["excel_file_id"]))
            if deleted is None:
                return None
        else:
            with self._connect() as connection:
                connection.execute("DELETE FROM archive_documents WHERE id = ?", (document_id,))
            self._document_path(row).unlink(missing_ok=True)
        return self._public_document(row)

    def reindex(self) -> int:
        rows = self._load_rows()
        indexed = 0
        for row in rows:
            if row["excel_file_id"] is not None:
                content = self._excel_content(int(row["excel_file_id"]))
            else:
                try:
                    content = _extract_document_text(
                        self._document_path(row).read_bytes(), f".{row['extension']}"
                    )
                    content = content[: self.MAX_INDEX_CHARS]
                except Exception as exc:
                    self._update_status(row["id"], "failed", str(exc))
                    continue
            status = "ready" if content else "empty"
            with self._connect() as connection:
                connection.execute(
                    "UPDATE archive_documents SET content_text = ?, status = ?, error = NULL WHERE id = ?",
                    (content, status, row["id"]),
                )
            self._replace_fts(row["id"], content + " " + self._search_text(row))
            indexed += 1
        return indexed

    def preview_workbook(self, document_id: int) -> dict[str, Any]:
        row = self._get_row(document_id)
        if row is None or row["kind"] != "bom":
            raise ArchiveError("BOM not found")
        path = self._document_path(row)
        if row["extension"] == "xls":
            import xlrd

            workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
            sheets = []
            for sheet in workbook.sheets():
                rows = [
                    [str(value).strip() for value in sheet.row_values(row_index)[: self.MAX_PREVIEW_COLUMNS]]
                    for row_index in range(min(sheet.nrows, self.MAX_PREVIEW_ROWS))
                ]
                sheets.append({"name": sheet.name, "rows": rows})
            return {"title": row["title"], "sheets": sheets}

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in workbook.worksheets:
                rows = []
                for values in sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row or 0, self.MAX_PREVIEW_ROWS),
                    max_col=min(sheet.max_column or 0, self.MAX_PREVIEW_COLUMNS),
                    values_only=True,
                ):
                    rows.append(["" if value is None else str(value) for value in values])
                sheets.append({"name": sheet.title, "rows": rows})
            return {"title": row["title"], "sheets": sheets}
        finally:
            workbook.close()

    def document_path(self, document_id: int) -> tuple[dict[str, Any], Path] | None:
        row = self._get_row(document_id)
        if row is None:
            return None
        return row, self._document_path(row)

    def _migrate_existing_excel_files(self) -> None:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT id, original_name, stored_path, sha256, size_bytes, status, error, created_at
                FROM excel_files
                WHERE id NOT IN (SELECT COALESCE(excel_file_id, -1) FROM archive_documents)
                """
            ).fetchall()
        for row in rows:
            self.register_excel_file(dict(row))

    def _load_rows(self) -> list[dict[str, Any]]:
        with self._connect() as connection:
            return [dict(row) for row in connection.execute("SELECT * FROM archive_documents ORDER BY created_at DESC, id DESC")]

    def _get_row(self, document_id: int | None) -> dict[str, Any] | None:
        if document_id is None:
            return None
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM archive_documents WHERE id = ?", (document_id,)).fetchone()
        return dict(row) if row is not None else None

    def _get_row_by_hash(self, digest: str) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM archive_documents WHERE sha256 = ?", (digest,)).fetchone()
        return dict(row) if row is not None else None

    def _get_row_by_excel_id(self, excel_file_id: int) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM archive_documents WHERE excel_file_id = ?", (excel_file_id,)).fetchone()
        return dict(row) if row is not None else None

    def _excel_content(self, file_id: int) -> str:
        with self._connect() as connection:
            values = connection.execute(
                "SELECT sheet_name, raw_value FROM excel_cells WHERE file_id = ? ORDER BY sheet_id, row_number, column_number",
                (file_id,),
            ).fetchall()
        return " ".join(f"{row['sheet_name']} {row['raw_value']}" for row in values)

    def _matches(
        self,
        row: dict[str, Any],
        query: str,
        kind: str,
        category: str,
        vendor: str,
        board_code: str,
        main_chip: str,
    ) -> bool:
        if kind and row["kind"] != kind:
            return False
        if category and row["category"] != category:
            return False
        if vendor and row["vendor"] != vendor:
            return False
        if board_code and row["board_code"] != board_code:
            return False
        if main_chip and row["main_chip"] != main_chip:
            return False
        terms = _query_terms(query)
        haystack = _normalize_text(self._search_text(row))
        return not terms or all(term in haystack for term in terms)

    @staticmethod
    def _search_text(row: dict[str, Any]) -> str:
        return " ".join(
            str(row.get(key) or "")
            for key in (
                "title", "category", "package", "vendor", "remark", "note",
                "board_code", "board_type", "main_chip", "board_name",
                "original_name", "content_text",
            )
        )

    def _public_document(self, row: dict[str, Any] | None) -> dict[str, Any]:
        if row is None:
            raise ArchiveError("document metadata is missing")
        return {
            "id": row["id"],
            "title": row["title"],
            "kind": row["kind"],
            "category": row["category"] or "未分类",
            "package": row["package"],
            "vendor": row["vendor"],
            "remark": row["remark"],
            "note": row["note"],
            "board_code": row["board_code"],
            "board_type": row["board_type"],
            "main_chip": row["main_chip"],
            "board_name": row["board_name"],
            "original_name": row["original_name"],
            "extension": row["extension"],
            "mime_type": row["mime_type"],
            "size_bytes": row["size_bytes"],
            "size": _format_size(row["size_bytes"]),
            "status": row["status"],
            "error": row["error"],
            "created_at": row["created_at"],
            "file_modified_at": row["file_modified_at"],
        }

    def _document_path(self, row: dict[str, Any]) -> Path:
        path = (self.storage_dir / row["stored_name"]).resolve()
        root = self.storage_dir.resolve()
        if root != path and root not in path.parents:
            raise ArchiveError("document path is outside storage directory")
        return path

    def _replace_fts(self, document_id: int, content: str) -> None:
        with self._connect() as connection:
            connection.execute("DELETE FROM archive_documents_fts WHERE document_id = ?", (str(document_id),))
            connection.execute(
                "INSERT INTO archive_documents_fts (document_id, content) VALUES (?, ?)",
                (str(document_id), content),
            )

    def _update_status(self, document_id: int, status: str, error: str | None) -> None:
        with self._connect() as connection:
            connection.execute("UPDATE archive_documents SET status = ?, error = ? WHERE id = ?", (status, error, document_id))

    def _connect(self) -> sqlite3.Connection:
        connection = connect_database(self.db_file, "archive")
        connection.row_factory = sqlite3.Row
        return connection


def parse_bom_filename(filename: str) -> dict[str, str]:
    stem = Path(filename).stem.strip()
    identity = re.match(r"^(RD|PD)_([A-Z][A-Z0-9]*\d[A-Z0-9]*)_", stem, re.IGNORECASE)
    if identity is None:
        return {"board_code": "", "board_type": "", "main_chip": "", "board_name": ""}
    board_code = identity.group(1).upper()
    versioned = re.match(r"^(RD|PD)_([A-Z][A-Z0-9]*\d[A-Z0-9]*)_.*?_V\d+(?:\.\d+)+", stem, re.IGNORECASE)
    components_at = stem.upper().rfind("_COMPONENTS LIST")
    return {
        "board_code": board_code,
        "board_type": "开发板" if board_code == "RD" else "产品板",
        "main_chip": identity.group(2).upper(),
        "board_name": versioned.group(0) if versioned else stem[:components_at] if components_at >= 0 else "",
    }


def _extract_document_text(payload: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(payload))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = Document(io.BytesIO(payload))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        paragraphs.extend(" ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return "\n".join(paragraphs)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            return "\n".join(
                f"工作表：{sheet.title}\n" + "\n".join(
                    ",".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                for sheet in workbook.worksheets
            )
        finally:
            workbook.close()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=payload, on_demand=True)
        return "\n".join(
            f"工作表：{sheet.name}\n" + "\n".join(",".join(str(value) for value in sheet.row_values(index)) for index in range(sheet.nrows))
            for sheet in workbook.sheets()
        )
    return payload.decode("utf-8", errors="replace")


def _sha256(payload: bytes) -> str:
    import hashlib

    return hashlib.sha256(payload).hexdigest()


def _normalize_text(value: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value or "")).strip().casefold())


def _query_terms(query: str) -> list[str]:
    return [term for term in re.split(r"[\s,，;；/]+", _normalize_text(query)) if term]


def _text(value: Any) -> str:
    return str(value or "").strip()


def _excerpt(content: str, terms: list[str]) -> str:
    normalized = _normalize_text(content)
    if not normalized:
        return ""
    position = min((normalized.find(term) for term in terms if normalized.find(term) >= 0), default=0)
    start = max(0, position - 120)
    return normalized[start : start + 360]


def _format_size(size: int) -> str:
    value = float(size or 0)
    units = ("B", "KB", "MB", "GB")
    unit_index = 0
    while value >= 1024 and unit_index < len(units) - 1:
        value /= 1024
        unit_index += 1
    return f"{value:.0f} {units[unit_index]}" if unit_index == 0 else f"{value:.1f} {units[unit_index]}"


def mime_type(suffix: str) -> str:
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xls": "application/vnd.ms-excel",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
        ".xltx": "application/vnd.openxmlformats-officedocument.spreadsheetml.template",
        ".xltm": "application/vnd.ms-excel.template.macroEnabled.12",
        ".txt": "text/plain; charset=utf-8",
        ".md": "text/markdown; charset=utf-8",
    }.get(suffix.lower(), "application/octet-stream")
