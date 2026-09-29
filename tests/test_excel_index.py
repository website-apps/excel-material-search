import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from test_support import connect_sqlite
import backend.excel_index as excel_module
import backend.spec_archive as archive_module
excel_module.connect_database = connect_sqlite
archive_module.connect_database = connect_sqlite

from backend.excel_index import ExcelIndexError, ExcelIndexStore


def workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Parts"
    first["A1"] = "Parts inventory"
    first.append([])
    first.append(["Part Number", "Reference", "Description"])
    first.append(["WPM3401", "R1", "Power regulator"])
    first.append([None, "R2", None])
    first.append(["WPM3401", "R3", "Power regulator duplicate"])

    second = workbook.create_sheet("Second Sheet")
    second.append(["generated", "metadata"])
    second.append([])
    second.append([])
    second.append(["Description", "Part Number", "Designator"])
    second.append(["Power regulator", "WPM3401", "U1"])
    second.append(["Other", "WPM3402", "U2"])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def description_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM"
    sheet.append(["Part Number", "Description"])
    sheet.append(["C100", "CAP 22uF TANT 10V 20%"])
    sheet.append(["C101", "CAP 22uF X5R 6.3V 20%"])
    sheet.append(["C102", "CAP 10uF X5R 20A"])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ExcelIndexStoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.store = ExcelIndexStore(
            Path(self.tempdir.name) / "excel.sqlite3",
            Path(self.tempdir.name) / "files",
        )

    def test_indexes_all_sheets_and_returns_exact_matches_with_context(self) -> None:
        result = self.store.index_workbook("sample.xlsx", workbook_bytes())

        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["sheet_count"], 2)
        self.assertTrue(Path(result["stored_path"]).exists())

        matches = self.store.search("WPM3401")

        self.assertEqual(len(matches), 3)
        self.assertEqual({match["sheet_name"] for match in matches}, {"Parts", "Second Sheet"})
        self.assertEqual({match["field_name"] for match in matches}, {"Part Number"})
        self.assertTrue(all(match["row_number"] >= 1 for match in matches))
        self.assertTrue(all(match["row_context"] for match in matches))
        self.assertTrue(
            all(match["created_at"] == result["created_at"] for match in matches)
        )

    def test_excludes_reference_and_designator_columns_from_search(self) -> None:
        self.store.index_workbook("sample.xlsx", workbook_bytes())

        self.assertEqual(self.store.search("R1"), [])
        self.assertEqual(self.store.search("U1"), [])

    def test_normalizes_whitespace_and_case_for_exact_queries(self) -> None:
        self.store.index_workbook("sample.xlsx", workbook_bytes())

        matches = self.store.search("  wpm3401 ")

        self.assertEqual(len(matches), 3)

    def test_falls_back_to_contains_match_for_partial_description_queries(self) -> None:
        self.store.index_workbook("capacitors.xlsx", description_workbook_bytes())

        matches = self.store.search("CAP 22uF")

        self.assertEqual(len(matches), 2)
        self.assertEqual(
            {match["raw_value"] for match in matches},
            {
                "CAP 22uF TANT 10V 20%",
                "CAP 22uF X5R 6.3V 20%",
            },
        )

    def test_treats_like_wildcards_as_literal_query_text(self) -> None:
        self.store.index_workbook("capacitors.xlsx", description_workbook_bytes())

        percent_matches = self.store.search("20%")
        underscore_matches = self.store.search("20_")

        self.assertEqual(len(percent_matches), 2)
        self.assertEqual(underscore_matches, [])

    def test_deletes_workbook_and_cascades_index_records(self) -> None:
        indexed_file = self.store.index_workbook("sample.xlsx", workbook_bytes())
        stored_path = Path(indexed_file["stored_path"])

        deleted_file = self.store.delete_file(indexed_file["id"])

        self.assertEqual(deleted_file["id"], indexed_file["id"])
        self.assertFalse(stored_path.exists())
        self.assertEqual(self.store.list_files(), [])
        self.assertEqual(self.store.search("WPM3401"), [])
        self.assertIsNone(self.store.delete_file(indexed_file["id"]))

    def test_deletes_multiple_workbooks_and_cascades_all_index_records(self) -> None:
        first_file = self.store.index_workbook("sample.xlsx", workbook_bytes())
        second_file = self.store.index_workbook(
            "capacitors.xlsx", description_workbook_bytes()
        )

        deleted_files = self.store.delete_files([second_file["id"], first_file["id"]])

        self.assertEqual(
            {deleted_file["id"] for deleted_file in deleted_files},
            {first_file["id"], second_file["id"]},
        )
        self.assertEqual(self.store.list_files(), [])
        self.assertEqual(self.store.search("WPM3401"), [])
        self.assertEqual(self.store.search("CAP 22uF"), [])

    def test_rejects_invalid_workbook_bytes(self) -> None:
        with self.assertRaises(ExcelIndexError):
            self.store.index_workbook("broken.xlsx", b"not an excel file")


if __name__ == "__main__":
    unittest.main()
